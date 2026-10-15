"""Rendu DOCX via python-docx - facture et rapport editables."""
from __future__ import annotations
from datetime import datetime
from io import BytesIO
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    raise ImportError("pip install python-docx")


def _set_cell_bg(cell, hex_color: str):
    """Applique un fond de couleur a une cellule."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.replace("#", ""))
    tcPr.append(shd)


def _add_paragraph(doc, text: str, size=10, bold=False, color=None,
                    align=None, space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text or "")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
    p.paragraph_format.space_after = Pt(space_after)
    return p


def render_facture_docx(config: dict, data: dict, totals: dict,
                         output_path: Path = None) -> bytes:
    """Genere une facture DOCX editable."""
    doc = Document()

    # Marges A4
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    societe = config.get("societe", {})
    client = data.get("client", {})
    devise = config.get("tva", {}).get("devise", "EUR")

    # En-tete : nom + titre
    _add_paragraph(doc, societe.get("nom", ""), size=14, bold=True,
                    color="0F172A", space_after=2)
    _add_paragraph(doc, f"{societe.get('adresse', '')} - {societe.get('code_postal', '')} {societe.get('ville', '')}",
                    size=8, color="64748B", space_after=1)
    _add_paragraph(doc, f"SIRET {societe.get('siret', '')} - TVA {societe.get('tva_intracom', '')}",
                    size=8, color="64748B", space_after=8)

    # Titre FACTURE
    _add_paragraph(doc, "FACTURE", size=22, bold=True, color="0F172A",
                    align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _add_paragraph(doc, f"N\u00b0 {data.get('numero', 'FA-2026-0001')}",
                    size=10, color="64748B",
                    align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    # Client
    _add_paragraph(doc, "FACTURE A", size=8, bold=True, color="2563EB", space_after=2)
    _add_paragraph(doc, client.get("nom", ""), size=11, bold=True, color="0F172A", space_after=1)
    if client.get("adresse"):
        _add_paragraph(doc, client.get("adresse", ""), size=9, color="334155", space_after=1)
    if client.get("code_postal"):
        _add_paragraph(doc, f"{client.get('code_postal', '')} {client.get('ville', '')}",
                        size=9, color="334155", space_after=8)

    # Dates
    _add_paragraph(doc, f"Date d'emission : {data.get('date', '')}",
                    size=9, color="334155", space_after=2)
    if data.get("date_echeance"):
        _add_paragraph(doc, f"Echeance : {data.get('date_echeance', '')}",
                        size=9, bold=True, color="2563EB", space_after=12)

    # Tableau des lignes
    lignes = totals.get("lignes", [])
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # En-tetes
    headers = ["DESCRIPTION", "QTE", "PRIX HT", "TVA", "TOTAL HT"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        _set_cell_bg(hdr_cells[i], "0F172A")
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for l in lignes:
        row = table.add_row().cells
        row[0].text = l.get("description", "")
        row[1].text = str(l.get("quantite", ""))
        row[2].text = l.get("prix_ht_fr", "")
        row[3].text = l.get("tva_pct_fr", "")
        row[4].text = l.get("ligne_ht_fr", "")
        for c in row:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)

    # Espace
    doc.add_paragraph()

    # Totaux
    _add_paragraph(doc, f"Sous-total HT : {totals.get('total_ht_fr', '')}",
                    size=10, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2)
    _add_paragraph(doc, f"TVA ({data.get('tva_pct', '20')}%) : {totals.get('total_tva_fr', '')}",
                    size=10, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2)
    _add_paragraph(doc, f"TOTAL TTC : {totals.get('total_ttc_fr', '')}",
                    size=14, bold=True, color="2563EB",
                    align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=12)

    # Mentions
    _add_paragraph(doc, "MENTIONS LEGALES", size=8, bold=True, color="0F172A", space_after=2)
    mentions = config.get("mentions", {})
    for key in ("conditions_paiement", "penalites_retard", "indemnite_recouvrement"):
        if mentions.get(key):
            _add_paragraph(doc, mentions.get(key), size=7, color="475569", space_after=1)

    if output_path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
